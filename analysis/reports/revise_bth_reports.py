"""Revise BTH reports so regional transition analysis, rather than corpus mechanics, leads."""
from __future__ import annotations

import argparse
import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib.colors import HexColor, white

from generate_bth_reports import (
    BLUE, GRAY, ORANGE, TEAL, FIG, HERE, OutlookPDF, add_body, add_references,
    load, refs_for, set_cell_shading, set_doc_defaults,
)


BLACK = RGBColor(0, 0, 0)


def academic_defaults(doc: Document) -> None:
    set_doc_defaults(doc)
    for style_name in ("Title", "Heading 1", "Heading 2", "Heading 3"):
        style = doc.styles[style_name]
        style.font.color.rgb = BLACK
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
    title_ppr = doc.styles["Title"]._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)
    doc.styles["Heading 1"].paragraph_format.space_before = Pt(14)
    doc.styles["Heading 1"].paragraph_format.space_after = Pt(8)
    doc.styles["Heading 2"].paragraph_format.space_before = Pt(9)
    doc.styles["Heading 2"].paragraph_format.space_after = Pt(5)


def academic_title(doc: Document, title: str, subtitle: str) -> None:
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(title)
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p2.add_run(subtitle)
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor.from_string("4B5563")


def add_compare_table(doc: Document) -> None:
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    widths = (1.05, 1.65, 2.15, 2.15)
    for i, (cell, text, width) in enumerate(zip(table.rows[0].cells, ("地区", "转型角色", "主要基础", "核心约束"), widths)):
        cell.width = Inches(width)
        cell.text = text
        set_cell_shading(cell, "D9E7F0")
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in cell.paragraphs[0].runs:
            run.bold = True; run.font.color.rgb = BLACK
    rows = [
        ("北京", "需求侧规则与创新应用中心", "服务经济、科研能力、公共治理和绿色金融", "能源资源外部依赖较高，超大城市减排空间分散"),
        ("天津", "港口工业与能源转换枢纽", "港口、石化、装备制造、汽车产业和综合能源项目", "工业与运输排放集中，存量设施更新周期较长"),
        ("河北", "产业深度改造与清洁能源供给基地", "钢铁建材、装备制造、风光资源和生态空间", "高碳产业比重大，转型成本与就业稳定相互牵制"),
    ]
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(values):
            cells[i].text = value
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if i == 0: cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            if row_index % 2 == 1: set_cell_shading(cells[i], "F4F7F9")


def build_overviews(records: list[dict]) -> Path:
    doc = Document()
    academic_defaults(doc)
    academic_title(doc, "京津冀绿色转型概述", "产业结构 能源系统与城市治理")
    sections = [
        ("北京", "北京市",
         "北京的绿色转型已从目标设定转向投资决策、城市运行和技术应用的系统重塑。市级碳达峰方案构成总体框架，固定资产投资项目节能审查和碳排放评价把能源与碳约束前移至项目立项阶段，可再生能源开发利用条例以及建设项目全过程监督制度，则把开发利用责任延伸至规划、建设、验收和运行环节。产业政策的重点不是扩大高耗能制造，而是依托科技创新、数字经济和专业服务能力，提高节能技术、绿色工厂、绿色供应链与碳管理服务的扩散效率。建筑、交通和公共机构是城市端减排的主要场景，既有建筑改造、综合能源站、公共领域车辆电动化和废旧动力电池回收正在形成相互连接的政策链条。区级实践进一步体现功能差异：城市副中心重视产业与城市建设同步降碳，经开区突出先进制造业和园区转型，西城侧重公共机构与绿色项目支持，密云等生态涵养区更强调生态价值保护与低碳发展结合。北京面临的关键矛盾，是本地能源资源有限而能源消费和高品质城市运行需求刚性较强，因此其转型成效不仅取决于本地节能，还取决于能否通过绿电交易、跨区域能源基础设施和统一碳核算规则，将外部清洁能源供给转化为可核验的消费侧减排。同时，数据中心、人工智能等新型用能需求增长，要求北京在创新产业扩张与能效约束之间建立更精细的动态平衡。"),
        ("天津", "天津市",
         "天津的绿色转型具有鲜明的港口工业城市特征，工业体系、港口物流和城市能源系统必须同步调整。市级碳达峰方案与工业领域碳达峰方案确定了能源结构优化、重点行业能效提升和资源循环利用的基本方向；固定资产投资项目节能审查与碳排放评价、温室气体重点排放单位核查等制度，则逐步把转型要求落实到企业和项目。能源方面，盐光互补、公共机构光伏、氢能和新能源汽车等应用，与天津既有的电力、石化和装备制造体系相互嵌合。交通方面，港口岸电、港内运输、重型货运和城市车辆更新共同决定港口经济的减排幅度，单纯扩大新能源车辆数量并不足以替代港航、铁路和公路之间的结构调整。工业方面，石化、钢铁、装备制造和汽车产业需要在设备更新、原料替代、过程电气化和资源循环利用之间形成连续改造路径。天津各区的转型条件差异明显：滨海新区承担港口和工业减排重任，津南等区域具备可再生能源项目基础，中心城区更侧重建筑、交通与公共机构节能。天津的优势在于产业链完整、港口场景集中，便于形成规模化示范；其难点在于存量设施资本强度高、更新周期长，能源价格、技术成熟度和产业竞争力会直接影响改造节奏。未来几年，天津能否把企业碳核查、港口能源数据和区级项目进度连接起来，将决定其能否从单项工程推进转向港产城整体转型。"),
        ("河北", "河北省",
         "河北的绿色转型首先是一场产业结构与能源供给方式的深度调整。钢铁、建材、化工和传统能源相关产业体量较大，使河北既承担较高的减排压力，也掌握京津冀产业链中不可替代的制造能力。工业领域碳达峰政策将节能降碳改造、绿色制造、资源综合利用和产业结构优化置于核心位置；对钢铁等行业而言，短期重点仍是能效提升、超低排放、余能利用和运输清洁化，中长期则取决于电炉钢、氢冶金、低碳原料和绿电供给的成本与规模。河北同时拥有张家口等可再生能源基地和广阔的新能源开发空间，其角色不仅是降低本地排放，也包括为京津提供稳定的清洁电力和调节能力。唐山、邯郸、沧州等工业城市需要处理高碳产业改造与就业、财政和产业链安全之间的关系；雄安新区可在绿色建筑、数字化能源管理和低碳城市基础设施方面形成增量示范；廊坊、保定等环京城市则与北京产业疏解、交通联系和空气质量协同高度相关。河北转型的难点在于成本和收益在区域间并不天然对称：清洁能源基地、生态保护和工业改造投资主要发生在河北，而部分能源消费和高附加值服务收益位于京津。由此，跨区域绿电交易、生态补偿、产业链协作和基础设施共建，不只是环境政策的补充，而是河北持续推进转型的重要经济条件。河北的转型速度最终取决于高碳产业能否形成具有市场竞争力的低碳产品，以及清洁能源优势能否转化为本地产业升级和长期收益。"),
    ]
    for i, (name, province, text) in enumerate(sections, 1):
        doc.add_paragraph(f"{i}  {name}", style="Heading 1")
        add_body(doc, text)
        add_references(doc, refs_for(records, province))
    out = HERE / "京津冀绿色转型三地概述.docx"
    doc.save(out)
    return out


def build_policy_report(records: list[dict]) -> Path:
    doc = Document()
    academic_defaults(doc)
    academic_title(doc, "京津冀绿色转型政策研究报告", "区域分工 协同机制与发展趋势")

    doc.add_paragraph("摘要", style="Heading 1")
    add_body(doc, "京津冀绿色转型并不是三个行政区域分别推进减排任务的简单叠加，而是能源供需、产业链、交通网络和生态环境在区域尺度上的重新组织。北京的比较优势在规则、技术、资本和应用场景，天津连接港口、工业与能源转换，河北承担重化工业深度改造、清洁能源供给和生态安全支撑。三地政策已经覆盖碳达峰、节能审查、工业绿色化、绿色建筑、绿色交通和生态治理，但跨区域协同仍主要集中于大气污染联防联控，面向碳排放、绿电消纳和产业链转型的制度连接尚不充分。未来数年，区域转型的主线将从单项技术推广转向能源系统、产业空间和基础设施协同；其成效取决于绿电供需能否匹配、重点产业能否形成低碳竞争力，以及转型成本与收益能否在三地之间合理分担。", False)

    doc.add_paragraph("一  研究背景", style="Heading 1")
    add_body(doc, "京津冀是我国经济活动和人口高度集聚的城市群，也是能源消费、工业生产、交通运输与生态环境问题相互交织的典型区域。北京以服务业、科技创新和总部经济为主，天津兼具港口、工业和城市功能，河北拥有规模较大的钢铁、建材、装备制造产业以及可再生能源和生态空间。三地产业结构差异决定了减排责任、技术路径和转型成本并不相同，却又通过电力、原材料、产品、交通和人口流动紧密连接。若只在行政区内部制定和评价政策，容易出现能源生产地与消费地分离、产业转移与排放转移混同、区域基础设施重复建设等问题。")
    add_body(doc, "近年来，三地绿色转型的政策基础逐步完善。碳达峰实施方案明确了总体方向，节能审查和碳排放评价开始影响新增项目，工业、建筑、交通和公共机构形成了更具体的行动安排，《京津冀美丽中国先行区建设行动方案》则把生态环境协同提升到区域层面。由此，研究重点已经从“是否转型”转向“以何种空间分工和制度连接完成转型”。本报告重点分析三地转型角色、关键部门变化及其相互依赖关系，政策文本仅作为识别制度方向和行动重点的证据。")

    doc.add_paragraph("二  三地形成互补而非同质的转型格局", style="Heading 1")
    add_body(doc, "京津冀内部不存在一条完全相同的转型路径。北京需要降低城市运行和消费端碳强度，并通过技术、金融和标准影响区域供给；天津需要在港口物流和工业体系中实现规模化改造；河北则必须在保证产业链稳定的同时降低重化工业排放，并扩大清洁能源供给。三者共同构成需求牵引、枢纽转换和供给重构的区域链条。")
    add_compare_table(doc)
    add_body(doc, "这种分工意味着三地不能仅以本地排放变化衡量政策效果。北京增加外购绿电可能降低消费侧排放，却要求河北配置新能源和调节资源；天津提高港口清洁运输比例，需要与河北工业货源和跨省交通网络同步；河北压减落后产能若缺少技术、资本和市场支持，则可能加大地方经济与就业压力。区域转型的实质，是在共同约束下重新配置能源、产业和基础设施。")

    doc.add_paragraph("三  能源系统正在从单向供给转向区域协同", style="Heading 1")
    doc.add_paragraph("一  北京的能源约束决定其必须扩大区域协作", style="Heading 2")
    add_body(doc, "北京本地可再生能源资源和大规模能源生产空间有限，建筑、交通、数据中心和公共服务构成主要用能场景。节能审查、碳排放评价和新能源电价机制有助于改善新增项目的能源表现，但深度减排仍依赖外部绿电、储能、需求响应和跨区域电网。北京的核心作用是形成稳定的绿色能源需求，并以长期采购、绿色金融和碳核算规则降低供给侧项目的不确定性。")
    doc.add_paragraph("二  天津的价值在于能源转换与综合应用", style="Heading 2")
    add_body(doc, "天津拥有港口、工业园区和大型能源基础设施，可在光伏、氢能、工业余能、储能和交通燃料替代之间形成综合应用。盐光互补和公共机构光伏体现了新能源场景扩展，港口岸电与新能源车辆则把能源转型延伸至物流系统。天津的挑战不是缺少单项技术，而是不同能源载体和行业数据尚未完全贯通，难以形成全生命周期的成本与碳效益比较。")
    doc.add_paragraph("三  河北决定区域清洁能源供给的上限", style="Heading 2")
    add_body(doc, "河北具备发展风电、光伏和新型储能的资源与空间条件，并承担向京津输送清洁能源的重要功能。随着新能源占比提高，电网调节、消纳能力和跨省交易机制的重要性将超过单纯装机规模。若新能源基地只承担发电功能，而设备制造、运维服务、储能和绿电收益没有在当地形成产业闭环，资源优势难以转化为稳定的区域发展动力。")

    doc.add_paragraph("四  产业转型呈现三条不同路径", style="Heading 1")
    add_body(doc, "北京、天津和河北的产业转型分别对应技术与服务赋能、存量工业系统改造和高碳基础产业重构。三条路径具有先后衔接关系：北京形成标准、技术和专业服务，天津在复杂工业与港口场景中验证规模化应用，河北则决定钢铁、建材等基础材料能否实现显著降碳。")
    doc.add_picture(str(FIG / "figure_02_topic_heatmap.png"), width=Inches(5.35))
    cap = doc.add_paragraph("图1  三地政策议题结构。比例反映政策关注方向，不代表实际转型绩效。")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.runs[0].font.size = Pt(8.5)
    add_body(doc, "政策议题结构表明，北京和天津的能源转型、综合绿色转型和碳管理工具较为多元；河北现有制度信号更集中于工业和生态环境治理。这种差异与三地经济结构相符，但也提示区域政策连接仍然不足。例如，钢铁和汽车产业链涉及河北原材料、天津制造与港口、北京研发和市场，若碳足迹规则、绿电使用和低碳材料采购彼此割裂，单个企业的改造难以转化为全链条竞争优势。")
    doc.add_paragraph("五  城市建设与交通转型构成共同应用场景", style="Heading 1")
    add_body(doc, "建筑和交通是三地政策能够形成直接协同的领域。北京在超低能耗建筑、既有建筑改造和公共领域车辆更新方面具有较强需求，天津可以依托港口、城市更新和综合能源项目形成规模化示范，河北则提供绿色建材、新能源装备和跨城运输网络。建筑领域的关键变化，是从设计阶段的节能标准转向运行阶段的实际能耗与碳排放管理；交通领域则从乘用车电动化进一步转向港口机械、重型货运、铁路运输和充换电网络协同。")
    add_body(doc, "三地交通联系频密，使运输方式调整具有明显的区域外部性。天津港和河北工业城市之间的货运组织、北京都市圈通勤以及机场港口集疏运，都会影响新能源车辆的实际减排效果。只有当车辆、能源补给、线路组织和货源结构同步变化，绿色交通才能从设备替换转化为运输系统效率提升。")

    doc.add_paragraph("六  生态环境协同正在向减污降碳一体化延伸", style="Heading 1")
    add_body(doc, "大气污染联防联控是京津冀最成熟的区域协同机制，为绿色转型提供了制度基础。过去的协同更多围绕污染物排放、重污染天气应对和跨界生态保护展开，未来需要进一步纳入碳排放、能源结构和产业链变化。污染物和温室气体在钢铁、燃煤、交通和工业过程等领域具有同源性，设备更新、能源替代和运输结构调整可以产生协同效益；但并非所有污染治理措施都会自动带来显著降碳，一些末端治理设施还可能增加能源消耗。因而，减污降碳协同需要从口号转化为具体项目的能源、污染物和碳排放综合评估。")

    doc.add_paragraph("七  区域协同的核心矛盾是成本与收益分布不对称", style="Heading 1")
    add_body(doc, "京津冀转型中的许多收益具有区域共享属性，而成本集中在具体地区和企业。河北建设新能源基地、实施生态保护和改造高碳产业，需要承担资本投入与就业调整压力；北京和天津则可能通过清洁能源消费、环境改善和高附加值服务获得较多收益。类似地，港口和运输系统改造由天津承担较多基础设施成本，但减排收益分布在整条产业链。若缺少稳定的绿电交易、生态补偿、绿色采购和产业协作机制，地方政府与企业会倾向于选择短期成本更低的方案。")
    add_body(doc, "因此，评价区域转型不能停留在三地排放或政策数量的横向比较，而应识别跨区域物质流、能源流和价值流。绿电从何处生产、低碳材料由谁采购、运输减排由谁投资、生态价值如何补偿，决定了协同机制能否长期运行。区域内的共同规则应首先服务于这些具体关系，而不是追求形式上的政策一致。")

    doc.add_paragraph("八  2027至2030年的主要趋势", style="Heading 1")
    trends = [
        ("能源转型将更重视消纳和灵活性", "新能源装机继续增长后，电网调节、储能、需求响应和跨区交易将成为约束。北京的用能负荷、天津的综合能源场景和河北的新能源基地需要更紧密匹配。"),
        ("工业政策将转向产品碳竞争力", "钢铁、汽车、装备和化工企业将更多面对产品碳足迹、绿色采购和国际供应链要求。低碳产品能否获得价格、订单或融资优势，将影响企业改造意愿。"),
        ("城市减排将从建设标准转向运行绩效", "绿色建筑、综合能源站、充换电网络和公共机构项目将更加重视实际运行数据。设计指标与真实能耗之间的差异会成为政策评估重点。"),
        ("区域协同将由环境治理扩展至产业与能源治理", "跨界大气治理仍是基础，但绿电交易、绿色运输、低碳材料和产业链碳管理将成为新的协同领域。"),
    ]
    for title, text in trends:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Pt(18)
        p.paragraph_format.first_line_indent = Pt(-18)
        r = p.add_run(title + "。")
        r.bold = True
        p.add_run(text)

    doc.add_paragraph("结论", style="Heading 1")
    add_body(doc, "京津冀绿色转型的决定性变量不是三地各自出台多少政策，而是能否把差异化优势组织成稳定的区域关系。北京需要以需求、规则和创新能力牵引低碳供给，天津需要把港口和工业场景转化为规模化技术应用，河北需要在产业深度改造和清洁能源供给中获得与其投入相匹配的长期收益。三地之间的能源、产业、交通和生态联系越能被准确计量并纳入市场与公共政策，区域转型就越可能在减排、产业竞争力和经济稳定之间形成可持续平衡。")

    all_refs, seen = [], set()
    for province in ("北京市", "天津市", "河北省"):
        for item in refs_for(records, province):
            if item["url"] not in seen:
                seen.add(item["url"]); all_refs.append(item)
    add_references(doc, all_refs)
    out = HERE / "京津冀绿色转型政策研究报告.docx"
    doc.save(out)
    return out


def build_outlook_pdf() -> Path:
    out = HERE / "京津冀绿色转型展望报告.pdf"
    p = OutlookPDF(out)
    # 1
    p.header("2027至2030年")
    y = p.title("京津冀绿色转型展望", 690, 28)
    y = p.text("能源系统 产业结构与区域协同", 44, y, 500, 15, 23, ORANGE)
    y = 505
    y = p.text("核心判断", 44, y, 500, 12, 19, BLUE)
    y = p.text("京津冀绿色转型将从分地区推进单项任务，转向能源供需、产业链和基础设施的区域协同。北京提供需求、规则和创新场景，天津连接港口工业与能源转换，河北承担产业深度改造、清洁能源供给和生态安全支撑。", 44, y - 6, 500, 11, 20)
    y = p.text("未来数年的关键不只是扩大新能源或增加绿色项目，而是让绿电、低碳材料、绿色运输和生态价值在三地之间形成可计量、可交易、可持续的关系。", 44, y - 18, 500, 11, 20)
    p.text("本报告依据京津冀近三年公开政策及现行政策框架作出情景判断。", 44, 95, 500, 8.5, 13, GRAY)
    p.finish_page()
    # 2 background and structure
    p.header("01 区域结构")
    y = p.title("三地差异决定了转型必须协同推进", 750, 19)
    y = p.text("北京、天津和河北分别处于区域价值链的需求端、转换枢纽和供给基础。三地路径不能同质化，也不能彼此分离：北京扩大绿电消费需要河北提供电源与调节能力，天津港口减排依赖河北货源和跨省交通组织，河北产业改造则需要京津的技术、资本和低碳产品市场。", 48, y, 495, 10.5, 18)
    nodes = [("河北", "清洁能源\n基础材料"), ("天津", "港口工业\n能源转换"), ("北京", "市场规则\n技术应用")]
    xs = [120, 300, 480]
    for i, ((head, body), x) in enumerate(zip(nodes, xs)):
        p.c.setFillColor(HexColor("#F1F6F8")); p.c.circle(x, 470, 65, fill=1, stroke=0)
        p.c.setFillColor(HexColor("#22577A")); p.c.setFont("Hei", 13); p.c.drawCentredString(x, 487, head)
        p.c.setFillColor(HexColor("#52606D")); p.c.setFont("STSong-Light", 9)
        for j, line in enumerate(body.split("\n")): p.c.drawCentredString(x, 462 - j * 15, line)
        if i < 2:
            p.c.setStrokeColor(HexColor("#2A9D8F")); p.c.setLineWidth(2); p.c.line(x + 66, 470, xs[i + 1] - 66, 470)
    y = 340
    for text in ("北京的重点是降低城市运行和消费端碳强度，并通过标准、金融和技术服务影响区域供给。",
                 "天津的重点是改造港口、石化、装备制造和交通系统，形成规模化综合能源场景。",
                 "河北的重点是重化工业深度减碳、清洁能源扩张，并处理转型成本与就业稳定。"):
        y = p.text(text, 52, y, 485, 10, 17, bullet="•") - 5
    p.finish_page()
    # 3 energy
    p.header("02 能源系统")
    y = p.title("新能源增长之后，消纳和灵活性成为约束", 750, 18)
    y = p.text("河北的风光资源、天津的综合能源设施与北京的稳定负荷，构成区域能源协同的物质基础。随着新能源比例上升，决定性因素将由装机规模转向电网调节、储能、需求响应和跨区域交易。", 48, y, 495, 10.5, 18)
    stages = [("供给", "河北风电光伏与储能"), ("转换", "天津港口工业与综合能源"), ("消费", "北京建筑交通与新型负荷")]
    y0 = 500
    for i, (head, body) in enumerate(stages):
        x = 55 + i * 172
        p.c.setFillColor(HexColor("#EFF6F8")); p.c.roundRect(x, y0, 145, 90, 8, fill=1, stroke=0)
        p.c.setFillColor(HexColor("#22577A")); p.c.setFont("Hei", 11); p.c.drawString(x + 12, y0 + 61, head)
        p.text(body, x + 12, y0 + 38, 120, 8.8, 14)
    y = 440
    y = p.text("北京的新型数据中心和电动交通可能推高电力需求；天津存量工业设施的电气化会改变负荷曲线；河北新能源基地则面临调峰和外送能力约束。这三类变化必须在同一能源系统中观察。", 48, y, 495, 10, 18)
    p.text("观察重点：跨区绿电交易、储能利用率、需求响应能力、工业电气化进度以及新能源消纳水平。", 48, 170, 495, 9.5, 16, ORANGE)
    p.finish_page()
    # 4 industry
    p.header("03 产业转型")
    y = p.title("低碳竞争力将沿产业链传导", 750, 19)
    y = p.text("京津冀拥有钢铁、汽车、装备制造、石化和数字服务等完整产业链。未来的减排压力将不只作用于单个企业，而会通过产品碳足迹、绿色采购、融资条件和供应链准入沿产业链传导。", 48, y, 495, 10.5, 18)
    p.image(FIG / "figure_02_topic_heatmap.png", 92, 265, 410, 400)
    y = 235
    y = p.text("图中反映的是政策关注结构：北京和天津的议题较为多元，河北更集中于工业和生态治理。它说明三地转型角色不同，但不能用于评价实际绩效。", 48, y, 495, 9.5, 16, GRAY)
    p.text("真正的协同机会在产业链交界处：河北低碳钢材和绿色电力、天津汽车及装备制造、北京技术标准与采购市场可以共同塑造低碳产品需求。", 48, 145, 495, 10, 17, ORANGE)
    p.finish_page()
    # 5 Beijing
    p.header("04 北京")
    y = p.title("北京的作用是形成规则与应用牵引", 750, 19)
    paragraphs = [
        "北京已将节能与碳约束前移到固定资产投资和项目决策，并在建筑、交通、公共机构、绿色制造和资源循环等场景推进技术应用。其优势是科研、金融和专业服务能力强，可以降低新技术进入市场的制度成本。",
        "北京的能源资源外部依赖较高，本地节能与外部清洁能源采购必须同时推进。绿电消费只有与可追溯的电力来源、稳定的交易安排和统一碳核算相结合，才能形成真实、可核验的减排。",
        "人工智能和数据中心等新型负荷快速增长，使创新产业扩张与能源约束之间出现新的张力。未来北京的转型质量将更多取决于运行能效和系统调节能力，而不是示范项目数量。",
    ]
    for text in paragraphs:
        y = p.text(text, 48, y, 495, 10.5, 19) - 18
    p.text("重点变量：项目碳评价、绿电采购、既有建筑运行绩效、公共交通电动化和新型负荷能效。", 48, 165, 495, 9.5, 16, ORANGE)
    p.finish_page()
    # 6 Tianjin
    p.header("05 天津")
    y = p.title("天津的关键在港口与工业系统整体改造", 750, 19)
    for text in (
        "天津的港口、石化、装备制造和汽车产业形成高强度、连续性的能源需求，也提供了新能源、氢能、工业余能和绿色物流规模化应用的条件。单项技术示范只有进入港产城系统，才能产生稳定减排。",
        "港口岸电、新能源车辆和清洁运输需要与货源组织、铁路集疏运和工业企业物流计划协同。若仅替换车辆而不改变运输结构，减排空间将受到限制。",
        "天津各区的产业和城市功能差异较大。滨海新区承担工业与港口减排，中心城区侧重建筑交通，津南等区域具备可再生能源项目基础。区级行动能否与企业核查和市级能源系统连接，是从项目改造走向城市转型的关键。",
    ):
        y = p.text(text, 48, y, 495, 10.5, 19) - 18
    p.text("重点变量：重点排放单位能效、港口清洁运输比例、工业设备更新周期、区级项目进度和能源基础设施协同。", 48, 150, 495, 9.5, 16, ORANGE)
    p.finish_page()
    # 7 Hebei
    p.header("06 河北")
    y = p.title("河北需要把减排压力转化为产业升级能力", 750, 19)
    for text in (
        "河北的钢铁、建材和化工产业决定了区域工业减排的深度。短期路径仍以能效提升、超低排放、余能利用和清洁运输为主，中长期则取决于电炉钢、氢冶金、低碳原料和绿电的经济性。",
        "张家口等地的风光资源使河北成为区域清洁能源供给基础。新能源规模扩大后，电网、储能、外送与本地消纳能力将决定资源价值，单纯增加装机不足以形成长期竞争优势。",
        "河北承担的工业改造、生态保护和能源建设成本较高，而部分消费与高附加值收益位于京津。绿电交易、生态补偿、低碳材料采购和产业链协作，直接影响河北能否在转型中形成稳定收益与就业。",
    ):
        y = p.text(text, 48, y, 495, 10.5, 19) - 18
    p.text("重点变量：单位产品碳排放、低碳材料市场需求、绿电消纳、清洁运输和重点工业城市就业结构。", 48, 150, 495, 9.5, 16, ORANGE)
    p.finish_page()
    # 8 outlook
    p.header("07 发展展望")
    y = p.title("区域协同将从污染治理延伸到能源与产业治理", 750, 18)
    stages = [
        ("能源", "从新能源装机转向消纳、储能、需求响应和跨区交易"),
        ("产业", "从企业节能转向产品碳足迹、低碳材料和供应链竞争力"),
        ("城市", "从建设标准转向建筑、交通和公共设施的运行绩效"),
        ("区域", "从大气联防联控转向能源、运输、产业链和生态价值协同"),
    ]
    y = 635
    for head, body in stages:
        p.c.setFillColor(HexColor("#22577A")); p.c.circle(75, y + 2, 22, fill=1, stroke=0)
        p.c.setFillColor(white); p.c.setFont("Hei", 9); p.c.drawCentredString(75, y - 2, head)
        p.text(body, 115, y + 7, 400, 10, 17)
        y -= 105
    p.c.setFillColor(HexColor("#F4F7F8")); p.c.roundRect(45, 95, 505, 120, 8, fill=1, stroke=0)
    p.text("总体判断", 58, 185, 470, 11, 17, BLUE)
    p.text("京津冀能否形成持续转型，取决于三地能否把差异化优势组织成稳定关系：北京形成需求和规则，天津完成复杂场景的规模化转换，河北在产业改造和清洁能源供给中获得长期收益。", 58, 158, 470, 9.5, 16)
    p.finish_page()
    p.save()
    return out


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--only", choices=("docx", "pdf", "all"), default="all")
    args = parser.parse_args()
    records, _ = load()
    outputs = []
    if args.only in {"docx", "all"}:
        outputs.extend([build_overviews(records), build_policy_report(records)])
    if args.only in {"pdf", "all"}:
        outputs.append(build_outlook_pdf())
    print(json.dumps({"outputs": [str(x) for x in outputs]}, ensure_ascii=False))


if __name__ == "__main__":
    main()
