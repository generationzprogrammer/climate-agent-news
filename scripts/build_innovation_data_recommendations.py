"""Build the October 2026 field-design report from a reproducible local audit."""
import json
from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'analysis/innovation_quality_20261023'
Q = json.loads((OUT/'substance_audit.json').read_text(encoding='utf-8'))
D = Document()
sec=D.sections[0]
sec.page_width=Cm(21);sec.page_height=Cm(29.7)
sec.top_margin=Cm(1.8);sec.bottom_margin=Cm(1.8);sec.left_margin=Cm(2);sec.right_margin=Cm(2)
for name in ['Normal','Title','Heading 1','Heading 2','Caption']:
    s=D.styles[name];s.font.name='Times New Roman';s.font.size=Pt(10.5);s.font.color.rgb=RGBColor(0,0,0)
    s.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Microsoft YaHei')
for s in D.styles:
    for e in list(s.element.xpath('.//w:pBdr')):e.getparent().remove(e)
D.core_properties.author='';D.core_properties.title='科产融合全球案例库数据采集与模式识别建议'
f=sec.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.CENTER
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');f._p.append(fld)
def font(r,size=10.5,bold=False):
    r.font.name='Times New Roman';r.font.size=Pt(size);r.font.bold=bold;r.font.color.rgb=RGBColor(0,0,0)
    rf=r._element.get_or_add_rPr().get_or_add_rFonts()
    for k,v in [('ascii','Times New Roman'),('hAnsi','Times New Roman'),('eastAsia','Microsoft YaHei')]:rf.set(qn('w:'+k),v)
    la=OxmlElement('w:lang');la.set(qn('w:val'),'en-US');la.set(qn('w:eastAsia'),'zh-CN');r._element.rPr.append(la)
def p(t,size=10.5):
    x=D.add_paragraph();x.paragraph_format.line_spacing=1.13;x.paragraph_format.space_after=Pt(7);font(x.add_run(t),size);return x
def h(t,level=1):
    x=D.add_paragraph(style='Heading '+str(level));x.paragraph_format.space_before=Pt(9);x.paragraph_format.space_after=Pt(7);font(x.add_run(t),14 if level==1 else 11.5,True)
def page(t):D.add_page_break();h(t)
def table(headers,rows,widths):
    t=D.add_table(rows=1,cols=len(headers));t.autofit=False
    for c,w in zip(t.columns,widths):c.width=Cm(w)
    pr=t._tbl.tblPr;b=OxmlElement('w:tblBorders')
    for key in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+key);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');b.append(e)
    pr.append(b)
    margins=OxmlElement('w:tblCellMar')
    for key in ['top','bottom','left','right']:
        e=OxmlElement('w:'+key);e.set(qn('w:w'),'85');e.set(qn('w:type'),'dxa');margins.append(e)
    pr.append(margins)
    for ri,row in enumerate([headers]+rows):
        cells=t.rows[0].cells if ri==0 else t.add_row().cells
        for c,txt,w in zip(cells,row,widths):
            c.width=Cm(w);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            para=c.paragraphs[0];para.paragraph_format.line_spacing=1.04;para.paragraph_format.space_after=Pt(2)
            font(para.add_run(str(txt)),9,ri==0)
            if ri==0:
                sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E9ECEF');c._tc.get_or_add_tcPr().append(sh)
        t.rows[ri]._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    p('',3)

x=D.add_paragraph(style='Title');font(x.add_run('科产融合全球案例库\n数据采集与模式识别建议'),19,True)
p('2026年10月23日',10)
h('一  数据建设应转向可核验的合作过程')
p('建议保留现有全球案例目录，把下一阶段的主要投入用于补充案例特有的合作任务、主体分工、资源交换、知识产权安排和实施结果。只有掌握谁向谁提供了什么、由谁作出决定、何时形成了何种结果，才能解释不同科产融合模式的差异。继续增加机构名称和计划介绍，有助于发现研究对象，却不能相应提高模式比较的解释力。')
p(f'对当前发布仓库的{Q["cases"]:,}条记录进行检查，其中机构或平台541条、项目672条，案例标识无重复。问题主要在内容层：444条名录案例使用相同的“约束”说明与相同的“可借鉴做法”；434条记录的“协同机制”落入完全相同的文本组。摘要全文虽不重复，但替换名称、地点后仍可能重复相同句式。因此，全文去重和字段非空检查不足以保证实质性信息质量。')
table(['检查结果','对研究的影响','处理建议'],[
['444条通用分析重复\n占全部记录36.6%','共同规则被误读为独立案例认识','计划规则单独存储；案例页只突出特有事实'],
['815条仅关联1条来源记录\n占67.2%','能够确认存在，不一定能证实运行效果','按断言补充来源，优先查合同、年报和评价报告'],
['384条起始年份未核验\n占31.7%','不能据此绘制完整历史发展曲线','保留未知；区分成立、认定、资助和实际运行日期'],
['53个高校基地缺少个体机制描述','认定身份不能代表独特转化机制','回访高校年报、制度文件和具体转化项目']
],[4.3,6.1,6.6])
p('这些比例以2026年10月19日导出的1,213条案例为分母，检查日期为10月23日。重复文字不等于案例无效，单一来源也不必然不可靠；上述结果用于确定补采优先顺序，不用于给机构评级。检查脚本及快照指纹与本报告一并保存在本地工作目录。',9)
h('建议形成两层数据产品',2)
p('目录层用于全球检索，允许只有身份、地域、领域和官方来源。深描层用于模式识别，要求合作关系、治理安排和至少一项有时间与证据的实施事实。两层共用案例标识，不复制计数，也不删除证据尚薄的既有记录。')

page('二  采集对象与核心字段')
p('以一次可识别的合作安排为案例基本单位，并保留层级关系。资助计划、承载机构、具体项目、合作事件和成果分别建表。例如，一个计划支持十个中心，每个中心承担多个项目，应建立关联，不能把所有层级相加解释为十余种独立模式。区域生态是分析对象，需要由多机构、多项目及地区统计共同刻画。')
table(['对象','建议字段及类型','采集要求'],[
['案例身份','case_id 文本；official_id 文本；name_original、name_zh 文本；record_type 枚举；parent_id 关联','稳定标识；原名不被译名覆盖；区分计划、机构、项目'],
['空间位置','focus_country 国别代码；admin_code 行政区划码；site_name 文本；location_role 枚举','总部、实际实施地和合作方所在地分别记录；行政区划带版本'],
['时间状态','event_type 枚举；event_date 日期；date_precision 枚举；valid_from、valid_to 日期','成立、入选、启动、结项分别记；网页更新时间不得代替事件时间'],
['技术与需求','technology_tags 多标签；problem_statement 文本；target_user 文本；application 场景','说明解决哪一项技术或生产问题、服务谁；不要只写新能源等大类'],
['合作主体','organisation_id；name；type；country；role；effective_period','一行一主体及其角色；企业、大学、研究所、公共机构分别标识'],
['合作关系','from_org；to_org；flow_type；resource；agreement_type；date','一行一项交换或合作关系；明确资金、人才、技术、数据或设施流向'],
['治理安排','agenda_setter；decision_body；selection_rule；ip_rule；access_rule','谁提出课题、选择项目、持有知识产权；说明准入与使用条件'],
['资金与投入','amount 数值；currency 币种；price_year 年；funding_kind；period；payer','承诺、批准、拨付、支出分开；不得把多年总额当作年投入'],
['实施与结果','result_type；measure；value；unit；baseline；time_window；status','目标、活动、产出、实际采用和影响分开；数字必须带口径'],
['逐项证据','claim_id；field；source_id；locator；short_excerpt；review_status','每项核心事实定位至具体页面、表格或段落；规则与个案证据分开']
],[2.3,8.2,6.5])
p('采用长表存储可重复出现的资金、人员、里程碑和成果，而非不断增加“项目一”“项目二”列。中文用于阅读，原文与受控代码用于核验和计算。机构消歧优先采用公开机构标识或官方注册编号，无法确定时保留候选关联，避免强行合并。')

page('三  哪些特征能够区分合作机制')
p('模式应由合作如何运行来识别，而不能由国家名称、机构称号或采集网站直接决定。要素流动受阻、生态适配错位等仍可作为问题标签，但需要具体事件支持，不能作为所有案例的默认属性。以下特征建议优先补采，均允许“未公开”或“待核验”。')
table(['特征','可观察字段与参数','适合的原始材料'],[
['需求由谁提出','企业委托、公共任务、研究者提议、共同协商；是否有需求书和验收人','征集指南、合同、项目任务书'],
['决策权如何分配','理事会成员类型、投票权、否决权、课题选择流程；记录生效期','章程、治理协议、年度报告'],
['资金如何组合','公共资助、会员费、企业现金、实物投入、股权投资；金额与用途','资助决定、经审计财报、预算说明'],
['知识产权怎样处置','背景与新增知识产权归属、独占或非独占许可、收益分配、发表审查期','标准协议、公开制度、许可合同摘要'],
['合作是否开放','加入条件、会员层级、使用费用、外国主体限制、数据与设施开放范围','入会规则、设施用户政策、数据许可'],
['由谁承担工程验证','中试设施所有者、操作者、用户；测试标准、规模、工况和验收结果','测试报告、中试验收、设施技术文件'],
['人才如何参与','联合聘任、产业博士、人员派驻、培训；人数、全时当量、周期与去向','合作协议、培养方案、就业或培训报告'],
['技术如何进入市场','许可、衍生企业、采购、供应链采用、共同生产；首单客户与时间','采购公告、许可记录、企业公告、产品认证'],
['合作是否持续','协议续期、成员进退、重复合作、停运或失败；原因和时间','年度成员名册、续约及终止公告'],
['制度如何跨境衔接','数据出境条件、出口管制、标准互认、争议解决、知识产权适用法','法规、条约、项目合规文件；保留有效期']
],[3.0,8.2,5.8])
p('上述参数不是越高越好。独占许可可能支持高风险产业投入，也可能限制扩散；成员增加可能扩大资源范围，也可能提高协调成本。应保留原始类别和条件，再检验其与结果的关系，避免先转换成统一的五分制。相关研究区分学术合作与商业化，并强调组织、关系和制度背景的差异[1–3]。')

page('四  建议形成的可比较指标')
p('指标建立在字段之上。先明确观察单位、分母和时间窗，再计算比率。下表是拟建设指标，不表示当前已具备相应数据。不同项目阶段与技术领域分别比较，项目目标不能进入实际绩效指标。')
table(['指标','计算口径','关键限制'],[
['企业投入比例','同一项目、同一期间企业实际现金投入 ÷ 实际总投入','实物投入另列；批准预算不能替代实际支出'],
['公共资金带动投入','同期间非公共实际投入 ÷ 公共实际投入','不作因果“撬动效应”解释；分母为零时不计算'],
['产业参与结构','已核实企业数 ÷ 同口径全部合作机构数；另列中小企业占比','名录不完整时只报列名数量，不估算比例'],
['合作持续性','观察期内续期合作关系数 ÷ 到期且可追踪关系数','失访与正式终止分开；尚未到期不入分母'],
['技术转化时滞','首次可核验产业采用日期 减 去技术披露或项目启动日期','两种起点分别报告；未转化项目属于右删失'],
['技术实施进展','每个已核验里程碑的日期、参数和工况；必要时记录技术成熟度','成熟度等级需有判据；计划达到与实际达到分开'],
['成果采用比例','约定跟踪窗内已获产业采用成果数 ÷ 同期可追踪成果数','以成果而非论文数为分母；同一成果不重复计数'],
['许可及转化收入','机构年度已实现许可、转让、股权退出收入分别列示','合同额、到账额与收益不混用；保留币种和财年'],
['共享设施利用','实际使用机时 ÷ 可用机时；外部用户占比另列','先规定停机维护处理方式；跨设施须同类比较'],
['人才与就业结果','已完成培养人数、进入产业人数、明确追踪期限的去向比例','实习、全职就业和联合聘任不合并'],
['技术环境效果','在相同功能单位和系统边界下的能耗、排放或资源消耗差异','保留基线、检测或生命周期评价方法及不确定性'],
['地区吸收能力','地区研发人员、研发强度、相关产业就业结构等年度背景变量','采用官方统计；不归为单一案例的产出或影响']
],[3.4,7.5,6.1])
p('跨币种比较先保留名义原值；确需换算时明确汇率、价格基年及用途。研究投入的实际资源规模可另行考虑购买力平价，融资与合同支付通常需要市场汇率，两者不能混用。人员数量与全时当量同样不能直接相加。无来源数值记为缺失，不以零代替。')

page('五  模式识别应采用可观察的判据')
p('建议维持八类机制作为研究编码，网页入口仍可合并为少量大类。同一案例允许多标签，并明确主要机制的判断依据。下列条件属于建议编码规则，不是现成的国际统一标准，也不是对全部既有案例重新验证后的结论。')
table(['机制类型','必要的识别证据','适用前提与主要风险'],[
['竞争前联合研究','多家产业成员与研究机构共同确定议题；存在成员参与或共享研究规则','共性技术需求与可信协调者；竞争者可能限制敏感知识共享'],
['任务导向研发','公开目标、资助方、里程碑、承担团队及考核安排可核验','任务能分解并允许试错；短期考核可能压缩探索空间'],
['专业技术转移','独立或专职转移职能；有评估、许可、撮合或收益分配机制','稳定技术供给、专业服务人才及清晰权属；仅挂牌不足以成立'],
['共享中试与验证','有实际设施、测试服务、外部用户准入及工程验证任务','产业需求足以支撑利用率；重资产投入与需求错配风险'],
['创业与风险资本','研究成果关联新企业；有股权融资、创业支持或商业化治理安排','创业人才、融资及退出渠道；不能仅凭“孵化器”名称分类'],
['产业集群协同','多个独立组织存在持续合作网络、专业分工与区域协调机制','供应链和技能基础；单一园区地址不能证明集群合作'],
['能力建设','科研训练、人才或组织能力提升有实施安排，并连接产业需求','吸收能力和长期投入；培养人数不能直接代表商业化成效'],
['规则与标准协调','共同规则、标准或互认安排有参与方、决策程序与实施记录','协调权威及参与激励；签署倡议不能替代实际互认']
],[3.4,7.4,6.2])
h('国家与地区模式是条件组合',2)
p('“美国模式”“德国模式”可作叙述入口，不宜直接作为案例标签或先验答案。应先在相同观察层级内识别机制，再比较各地的资金、人才、产业需求和制度条件。例如，广东与浙江的差异，应通过相近产业、相近规模和相近时期的案例比较，并结合地区统计检验；不能用两个知名平台分别代表全省。')
p('来源计划可以提供共同规则，但不能据此认定每个成员都已完整实施。对计划规则推定的标签单独标记，只有个案章程、合同或运行事实支持时，才升级为个案机制证据。缺乏依据时保留“机制待核验”，不强制归入某一模式。')

page('六  从分类走向有解释力的比较')
h('先保证比较对象相同',2)
p('项目与项目、机构与机构分别比较。首先限定技术领域、发展阶段和观察年份，再匹配研究规模、产业结构及地区条件。公共计划支持的中心与大型企业自身研发组织即使名称相似，也可能承担不同任务。没有统一观察窗口时，累计专利、累计融资和年度收入不宜放在同一排名中。')
h('相似案例检索应说明相似在哪里',2)
p('可试用适合混合类型字段的距离方法，将治理安排、资金结构、产业接口和技术阶段作为输入。数值使用预先公布的尺度，类别按一致或差异处理；缺失项不参与该项距离计算，也不被视为相同。页面同时呈现实际可比较的字段数，以及贡献最大的共同点与差异。初期宜采用透明规则与人工核验，待字段稳定后再尝试聚类或潜类分析。')
p('建议把至少四个可核验机制维度作为相似案例试点的入选条件，并在三个至五个维度的不同设置下检查结果是否稳定。这是待试验的项目管理阈值，不是统计显著性标准。国家、来源站点和既有模式标签不参与首次聚类，以免算法只是重新发现数据采集渠道。')
h('识别特征与检验影响应分两步',2)
p('第一步描述哪些机制经常共同出现、在哪些条件下采用。第二步才研究它们是否关联更快转化、更高采用率或更低失败率。第二步需要时间顺序、可比对照和未成功案例；公开成功案例目录本身不足以识别政策效果。回归分析至少处理技术领域、项目年龄、规模及来源批次差异；案例数很多但来自少数计划时，有效独立信息远少于名义样本量。')
h('防止模型学习到来源模板',2)
p('训练与验证应按计划、机构或来源批次划分，不能把同一计划下的近似文字随机分到两边。补充逐一剔除来源批次的敏感性检验，检查模式分布是否发生显著改变。模型输出应包含证据定位和反证条件，而非只给一个标签或分数。')
h('阶段变化应由事件支持',2)
p('中国五年规划期与国际十年期可以作为汇总窗口，底层仍保留实际事件日期。成立于“十二五”的机构可能在“十四五”改变资助或知识产权制度，两次观察应关联同一机构，不能按成立年份推定其长期模式不变。全球治理缺位等问题还需记录具体制度障碍、发生时间、受影响主体与应对结果。')

page('七  三类现有案例的补采示范')
h('美国电源管理集成中心',2)
p('现有记录能够识别PMIC、2018年起始信息、Dartmouth与加州大学圣迭戈分校两家高校，以及Allegro、Meta、Richtek、Analog Devices、Marvell、Qualcomm、Enphase七个列名成员。网页可以直接展示这些可区分事实，无须重复“产学合作受到重视”。七个成员仅为来源列名，不自动等于当前全部成员或七家独立出资企业。')
p('下一步应补充具体研究课题、产业咨询委员会的选题权、会员费与研究权益、已完成的原型或测试结果。NSF的通用计划说明只能支持计划层面的竞争前研究安排，不能证明该中心某项技术已商业化[6]。')
h('澳大利亚低成本低排放熔融硫镍矿加工项目',2)
p('现有官方登记记录包含2025年9月至2027年9月的项目期、300万澳元资助、780.7121万澳元项目总价值，以及Oxleigh Recycling Technologies、Nico Resources、Curtin University等七个合作方。项目拟将用于电池回收的熔融硫技术扩展至红土镍矿加工，登记目标包括成本降低超过50%、用水降低85%、排放降低90%。这些是目标，不是已实现的性能。')
p('应优先补充中试规模、矿石品位、比较工艺、系统边界、连续运行时间、第三方测试和客户采用证据。没有这些条件，三个减幅数字不能与其他冶炼路线直接比较。该记录适合展示“目标—验证节点—实测结果”时间线，而非把预算与目标拼成成功案例[7]。')
h('中国高校科技成果转化基地',2)
p('71个教育部认定基地中，现有18个有额外的个体实践描述，其余53个主要反映认定与依托高校。例如东南大学记录涉及市场化转移组织、人员聘用、考核与激励，可以作为补采制度细节的线索；认定年份不能替代转移机构成立年份。')
p('应进一步取得制度文件的生效时间、收益分配条款、专业人员配置，以及按年度统计的许可、转让和到账收入。若只找到挂牌新闻，保留目录记录，不编写“显著提升转化效率”。若有具体许可项目，应单独建项目记录并关联基地，避免把个别交易外推为整个高校的持续绩效[8]。')
p('以上示范基于现有官方来源链接与登记字段。涉及机构现行成员、最新执行进展和实际绩效，须另行回源核验；本报告没有将这些待补信息写成已完成事实。',9)
p('本次网页修订另对PMIC、CITES、WindSTAR、BEST、GRAPES和CEPS六个中心补充了专属官方详情与十二项技术任务或设施事实。例如GRAPES资助说明列有6 MVA试验设施及15 kV测试能力，可用于识别工程验证条件，但不能据此计算技术转化成功率。新增内容与原名录字段并存，作为后续逐项补采的示范[10]。',10)

page('八  采集与验证工作的优先次序')
table(['阶段','应完成的工作','验收内容'],[
['先修正呈现和数据口径','分开计划规则与个案事实；统一目标、预算、观察值及缺失状态；保留旧记录','不再重复展示通用分析；旧标识与事实字段不丢失'],
['补采高价值细节','先追踪已有专属详情页或概要文件的案例，再补章程、年报、测试和转化文件','每个深描案例可回答技术任务、主体分工、实施事实；关键机制可定位'],
['建立分层核验集','按地区、机制、来源、语言及观察层级抽样；加入失败和中止案例','至少两名独立编码者先分别判断，再记录分歧与裁决'],
['最后开展模型比较','在核验集上检验分类与相似检索，再扩展到全库；保留版本与回滚','报告字段级准确率、召回率及类别混淆；区分人工核验与程序检查']
],[3.0,8.0,6.0])
p('建议先以约120条案例开展深描试点，而非立即为全库补写长文。该规模用于检查字段可获得性和工作负担，并非代表性样本量保证。重点覆盖中美，同时纳入其他地区、少数机制及不同成熟度；罕见类型可全部复核，抽样分析时保留入样概率或单列结果。成熟案例追踪成果，新近案例追踪里程碑，避免因时间不足把新项目判为失败。')
h('Agent负责提取候选事实而不是补齐空白',2)
p('采集Agent先发现官方详情页，提取Agent只从允许访问的材料抽取字段与短证据片段，核验步骤检查实体、日期、单位和事实状态，再由审核者处理制度解释、因果表述及冲突。保留模型版本、提示版本、来源哈希和字段变更记录。对缺失字段输出明确状态，不让模型根据国家、项目名称或同计划其他成员推断个案事实。')
p('缺失状态建议分为未公开、未采集、不适用、来源不可访问和存在冲突。首次补采未找到数据，不等于不存在合作或成果。自动检查重点包括金额与币种、时间逻辑、机构消歧、同一事实重复计数、事实来源定位，以及同批摘要中去除实体名后的重复句式；语义相似度只触发复核，不自动删除。')
h('定期更新以新增事实为目标',2)
p('每周检查近期项目状态与来源变更，每季度复核成员、年报及成果；已有页面不变时复用缓存。优先队列可综合最新事件、缺失关键机制、可访问专属材料和地域空白，而不以每周必须新增多少条为目标。保存新增事实数、实质修订数、核验积压及失效来源，不把重复改写计作更新。')

page('九  网页应呈现什么以及如何服务研究')
p('Metroverse值得借鉴的是统一对象口径、围绕具体问题组织比较，以及让用户沿关键指标深入查看。其公开说明强调城市与产业定义的可比性，也承认来源覆盖差异[4]。本库宜借鉴这一思路，而非直接移植城市经济复杂度指标；现有案例的采样方式与企业就业结构数据不同。')
table(['使用问题','建议展示','数据前提'],[
['这个案例究竟做了什么','顶部放技术任务；下方呈现主体、实施地、项目期、资助和阶段','有证据的事实卡；没有数据时不补空洞段落'],
['合作怎样发生','参与者角色表与少量有方向的资源关系；点击查看证据','有角色与关系数据后才画网络，不靠名称共现连线'],
['和谁可以比较','同层级相似案例；列出共同点、关键差异与缺失维度','共同可观察字段足够；比较规则公开且可复算'],
['实施到了哪一步','日期明确的目标、测试、采用及终止事件时间线','计划与实际分层；不以网页发布日期充当成果日期'],
['哪种模式适合本地','机制条件矩阵，并对照地区人才、产业需求、设施和制度','先识别条件缺口，再提出政策选择；不按案例数量打分']
],[3.2,7.5,6.3])
p('案例页应将标题、实质摘要、结构化事实与来源入口置于主要位置。计划共同规则可以折叠查看，通用的“需要结合本地条件”等句子不再逐条出现。原文目标和技术参数允许展开，保持中英文对应。图表优先采用同口径比较表、成果时间线和条件矩阵；在缺少可比数值时，不绘制看似精确的综合雷达评分。')
h('对研究与政策应用的价值',2)
p('研究者由此可以分析不同治理安排与技术转化路径的关系，构建有时间顺序的合作网络，并比较相近地区的制度选择。政府或智库则可以先识别本地缺少的是工程验证、专业转移人员、持续资金还是企业需求，再筛选具备相近前提的案例。建议应写成有证据的条件判断，不能从国外案例知名度直接推出本地应复制的组织形式。')
p('面向Scientific Data，优先完善可复用的字段定义、逐项来源、抽样与排除记录、独立语义核验和版本化发布。期刊要求技术验证支持数据质量，字段填满与程序测试通过并不能替代内容核验[5]。正式分析还需说明成功案例披露倾向、来源批次集中和时间覆盖不足，不将便利采样描述为全球总体普查。')

page('参考文献与核验依据')
refs=[
'OECD, Eurostat. Oslo Manual 2018: Guidelines for Collecting, Reporting and Using Data on Innovation. 4th ed. OECD Publishing, 2018. 第6章知识流动与第10章对象法。https://doi.org/10.1787/9789264304604-en',
'Perkmann M, Salandra R, Tartari V, McKelvey M, Hughes A. Academic engagement: A review of the literature 2011–2019. Research Policy, 2021, 50(1):104114. https://doi.org/10.1016/j.respol.2020.104114',
'Boschma R A. Proximity and Innovation: A Critical Assessment. Regional Studies, 2005, 39(1):61–74. https://doi.org/10.1080/0034340052000320887',
'Harvard Growth Lab. Metroverse FAQ and About. https://metroverse.hks.harvard.edu/faq ; https://metroverse.hks.harvard.edu/about/team 。示例页面 https://metroverse.hks.harvard.edu/city/10687/similar-cities 。本报告借鉴其对象可比与问题导向设计，不复用其指标或算法。',
'Scientific Data. Submission Guidelines; For Referees. https://www.nature.com/sdata/submission-guidelines ; https://www.nature.com/sdata/policies/for-referees 。访问日期2026-10-23。',
'U.S. National Science Foundation. IUCRC University information; Power Management Integration Center directory entry. https://iucrc.nsf.gov/universities/ ; https://iucrc.nsf.gov/centers/integrated-power-management-circuits-and-systems---power-one-ic',
'Australian Government. CRC Projects selection round outcomes. https://business.gov.au/grants-and-programs/cooperative-research-centres-projects-crcp-grants/crc-projects-selection-round-outcomes 。项目题名 Transforming Nickel Ore Processing with Low-Cost Low-Emission Molten Sulfur。',
'中华人民共和国教育部. 高校科技成果转化和技术转移基地认定及首批基地典型经验相关材料. https://www.moe.gov.cn/s78/A16/tongzhi/202004/t20200417_444200.html 。认定与具体实践分别取证。',
'科产融合全球案例库. 发布仓库提交a95411c的静态数据 innovation_cases.json，数据标记更新日期2026-10-19，核验日期2026-10-23。Windows工作副本SHA-256：'+Q['snapshot_sha256']+'。复核程序 scripts/audit_innovation_substance.py；基线与Notebook保存于 analysis/innovation_quality_20261023。',
'U.S. National Science Foundation. Center for Grid-Connected Advanced Power Electronic Systems. Award 1939124, facilities description. https://iucrc.nsf.gov/centers/grid-connected-advanced-power-electronic-systems/ 。其余五例的来源与章节定位保存在 config/innovation_case_details.json。'
]
for i,s in enumerate(refs,1):p(f'[{i}] {s}',9)
p('指标与模式判据为本报告面向该项目提出的设计建议。文献用于界定概念和测量原则，不代表文献已经验证本库的标签、阈值或推断。案例示范使用现有登记材料，正式研究引用时仍应保留具体来源版本与字段证据。',9)
OUT.mkdir(parents=True,exist_ok=True)
dest=OUT/'科产融合全球案例库数据采集与模式识别建议.docx'
D.save(dest)
print(dest)
