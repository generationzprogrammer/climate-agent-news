"""Generate reviewable, source-backed Chinese outreach documents (no API calls)."""
from pathlib import Path
import json
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'analysis/reports/platform_publication_20261011'
QA = ROOT / 'tmp/platform-publication-qa'
SITE = 'https://generationzprogrammer.github.io/climate-agent-news/'
DATA = json.loads((ROOT / 'static/data/innovation_cases.json').read_text(encoding='utf-8'))
ENERGY = json.loads((ROOT / 'static/data/bth_energy_database.json').read_text(encoding='utf-8'))
OUT.mkdir(parents=True, exist_ok=True)
QA.mkdir(parents=True, exist_ok=True)

def font(style, size, chinese='宋体', bold=False):
    style.font.name = 'Times New Roman'
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), chinese)

def document(title, subtitle):
    doc = Document()
    for style in doc.styles:
        if style.element.pPr is not None:
            for border in list(style.element.pPr.findall(qn('w:pBdr'))):
                style.element.pPr.remove(border)
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(.75)
    section.left_margin = section.right_margin = Inches(.85)
    font(doc.styles['Normal'], 11)
    doc.styles['Normal'].paragraph_format.line_spacing = 1.25
    doc.styles['Normal'].paragraph_format.space_after = Pt(7)
    for name, size in [('Title', 21), ('Heading 1', 15), ('Heading 2', 12)]:
        font(doc.styles[name], size, '黑体', True)
        doc.styles[name].paragraph_format.space_before = Pt(12)
        doc.styles[name].paragraph_format.space_after = Pt(7)
    font(doc.styles['Subtitle'], 10)
    doc.styles['Subtitle'].font.italic = False
    header = section.header.paragraphs[0]
    header.text = '格润  ·  全球气候与能源信息智能分析支持系统'
    font(doc.styles['Header'], 9)
    footer = section.footer.paragraphs[0]
    footer.alignment = 2
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)
    doc.core_properties.author = '格润项目组'
    doc.core_properties.last_modified_by = '格润项目组'
    doc.core_properties.title = title
    doc.core_properties.subject = subtitle
    doc.add_paragraph(title, 'Title')
    doc.add_paragraph(subtitle, 'Subtitle')
    return doc

def p(doc, text):
    return doc.add_paragraph(text)

def h(doc, text):
    return doc.add_paragraph(text, 'Heading 1')

def url(doc, label, value):
    paragraph = doc.add_paragraph()
    run = paragraph.add_run(label + '：'); run.font.size = Pt(9)
    hyperlink = OxmlElement('w:hyperlink')
    rid = doc.part.relate_to(value, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
    hyperlink.set(qn('r:id'), rid)
    run_xml = OxmlElement('w:r'); props = OxmlElement('w:rPr')
    fonts = OxmlElement('w:rFonts')
    for key, val in [('ascii','Times New Roman'), ('hAnsi','Times New Roman'), ('eastAsia','宋体')]: fonts.set(qn('w:'+key), val)
    props.append(fonts)
    color = OxmlElement('w:color'); color.set(qn('w:val'), '000000'); props.append(color)
    size = OxmlElement('w:sz'); size.set(qn('w:val'), '18'); props.append(size)
    underline = OxmlElement('w:u'); underline.set(qn('w:val'), 'single'); props.append(underline)
    run_xml.append(props); text = OxmlElement('w:t'); text.text = value; run_xml.append(text)
    hyperlink.append(run_xml); paragraph._p.append(hyperlink)

def diagram(name, labels, captions):
    image = Image.new('RGB', (1560, 290), 'white')
    draw = ImageDraw.Draw(image)
    face = 'C:/Windows/Fonts/msyh.ttc'
    large, small = ImageFont.truetype(face, 36), ImageFont.truetype(face, 26)
    for index, (label, caption) in enumerate(zip(labels, captions)):
        x = 10 + index * 395
        draw.rectangle((x, 25, x+350, 235), fill='#f6f7f8', outline='#aeb7bd', width=2)
        draw.text((x+175, 74), label, font=large, fill='#111111', anchor='mm')
        for row, line in enumerate(caption.split('|')):
            draw.text((x+175, 132+row*35), line, font=small, fill='#333333', anchor='mm')
        if index < 3:
            draw.line((x+360,130,x+380,130), fill='#555555', width=3)
            draw.polygon([(x+380,130),(x+371,124),(x+371,136)], fill='#555555')
    path = QA / (name+'.png'); image.save(path)
    return path

stats = DATA['statistics']
assert stats['cases'] == len(DATA['cases']) and len(DATA['analysis']['provinces']) == 31
profiles = DATA['analysis']['profiles']
model_count = len(profiles)
doc = document('从科研合作到产业应用', '科产融合全球案例库  |  对外介绍推文  |  2026年10月11日')
p(doc, '一项技术如何从实验室走向产业？不同国家如何组织企业、高校、研究机构与公共部门协同？格润在能源技术模式下建设科产融合全球案例库，将分散的项目与平台资料整理为可检索、可比较、可追溯的研究入口。')
h(doc, '753个案例，连接全球实践与中国区域经验')
p(doc, f'当前收录{stats["cases"]}个案例、{stats["sources"]}条来源，涉及{stats["country_tags"]}个国家标签，中国案例覆盖内地31个省级地区，并形成{model_count}组国家与区域模式综述。案例既包括正式登记的联合研发项目，也包括应用研究机构、中试平台、产业协作网络及技术规则安排。读者可以从具体技术问题出发，查找参与主体、合作方式、项目时间和原始证据。')
doc.add_picture(str(diagram('case-workflow', ['查找案例','理解机制','对齐比较','研判适配'], ['国家与地区|技术与组织方式','参与主体|研发与转化接口','同类项目|模式与区域差异','要素条件|可借鉴的政府工具'])), width=Inches(6.6))
h(doc, '不只看案例名称，更看合作如何发生')
p(doc, '“查看案例”展开研发任务、参与主体、协同机制与已经记录的实践。以德国弗劳恩霍夫为例，官方合作指南区分合同研发、长期战略伙伴、创新集群与衍生企业等转化渠道；企业需求并非只在研发完成后才进入，而是能够通过委托与持续合作参与研发组织。')
p(doc, '中国区域案例提供另一类观察。浙江清华长三角研究院以省校合作组织研发，再用概念验证、中试平台和制造业创新中心衔接工程应用，涉及柔性电子、陶瓷材料和氢能装备。对地方政府和智库而言，值得比较的是这些组织接口如何补齐当地产业链的验证与放大能力。')
doc.add_page_break()
h(doc, '从案例比较走向模式比较')
p(doc, '先选择一个技术领域，再勾选两至三个案例进行对比，可并列查看任务、主体、时间、资金及机制。具备同口径登记数据的项目提供合作结构雷达图，展示跨境联结、产业参与、科研支撑、主体多样性与联合资助规模，并保留原始值和出处。雷达图不作为技术领先或商业成功的排名。')
p(doc, '“国家与区域模式”进一步比较科研创业网络、任务导向攻关、应用研究桥接、联合研发、开放验证等组织方式，说明其成立所需的产业需求、科研能力、资金和制度条件。中国省份之间的比较，可帮助读者识别制造业基础、公共试验设施与校地合作网络对模式适配的影响。')
h(doc, '让案例研究服务于开放创新决策')
p(doc, '对政府部门，案例库支持围绕产业共性难题寻找组织工具，而不是简单复制某个地区的标签；对研究机构，可以检索潜在协作平台、比较中试与成果转化机制；对企业，可以了解联合研发、公共验证及跨区域合作的具体入口。筛选结果可导出，便于专题研究、访谈准备和报告引用。')
p(doc, '网站每日将已有情报中的近期动态与案例实体关联，已核验的案例事实和原始证据持续保留；新增案例经复核后入库。案例数量反映当前收录范围，不是全球创新活动的总体排名。')
h(doc, '进入案例库')
p(doc, '打开格润网站，切换至“能源技术”，点击顶部“科产融合”。可在“案例检索”“国家与区域模式”“宏观对比”之间切换；建议先筛技术领域，再比较两个组织方式不同的案例。')
url(doc, '案例库', SITE+'?mode=energy#innovationCases')
url(doc, '弗劳恩霍夫合作指南', 'https://www.fraunhofer.de/en/quick-links/companies/collaboration-Fraunhofer.html')
url(doc, '浙江清华长三角研究院简介', 'https://www.tsinghua-zj.edu.cn/institute/introduce')
doc.save(OUT/'科产融合全球案例库_对外介绍推文.docx')

doc = document('从全球新闻到每日重点', '格润新闻资讯与推送  |  对外介绍推文  |  2026年10月11日')
p(doc, '气候治理、能源技术与企业活动往往分散在不同机构和媒体网站。格润将持续采集、中文编译、地图展示与档案检索结合起来，让读者用一张地图和一份简报掌握值得跟进的变化，并能随时回到原文核查。')
h(doc, '三个现场，对应三类信息需求')
p(doc, '全球气候现场关注气候政策、国际谈判、碳市场、气候风险与减排举措。今日队列呈现最新自然日的合格情报，本周视图帮助读者回看连续七天的事件。点击地图上的地点，可阅读中文标题、具体概括及来源链接。')
p(doc, '全球能源技术现场关注能源技术进展、重大项目及产业动态，连接新能源、储能、人工智能数据中心等专题。读者可以从一条事件继续检索相关技术报告和历史文本，而不必在新闻、报告和数据库之间反复寻找入口。')
p(doc, '全球企业情报现场聚焦能源巨头的重要项目、初创企业的业务与技术、企业之间的合作。企业数据库补充重点企业的财务披露、核心技术、发展方向与项目资料；财务数据保留财年、币种和原始来源。')
doc.add_picture(str(diagram('news-workflow', ['持续采集','编译核验','地图与简报','归档与周报'], ['机构与专业媒体|公开RSS与网页','中文标题与概括|时间地点与证据','今日与本周|原文入口与PDF','增量文本积累|每周邮件订阅'])), width=Inches(6.6))
h(doc, '中文概括承载信息，不替代原文')
p(doc, '每条重点情报提供原新闻标题的中文译文和一段有实质内容的概括，尽量交代谁采取了什么行动、发生在何时何地、涉及怎样的数字或变化。采编保留原始标题、发布日期和来源，不以泛化的主题描述代替新闻内容。地点没有明确依据时不强行标注，重复链接与重复事件经过归并处理。')
doc.add_page_break()
h(doc, '每天更新网页，每周发送邮件')
p(doc, '日常自动采编安排在北京时间06:30和18:30启动，完成采集、编译、质量检查后发布。网站更新不依赖用户的电脑持续开机。定时任务可能受队列和来源可用性影响，阅读时可核对页面所示信息日期。')
p(doc, '“下载今日简报”与当日重点情报对应，以PDF集中呈现标题和一段话概括，原文链接统一放在文末。邮件订阅目前发送每周气候情报周报，安排在每周一北京时间08:00启动；它不是每日新闻逐条邮件，也不等同于所有模块的每日推送。')
h(doc, '怎样订阅与退订')
p(doc, '在网页点击“订阅”，填写常用邮箱并提交。页面显示“订阅已登记”才表示服务已确认接收；若出现超时或未确认，请重试或联系管理员，不能仅凭点击按钮判断订阅完成。每周邮件按已登记名单发送；未收到时先检查垃圾邮件。需要退订时，在同一窗口输入订阅邮箱，点击“退订”。')
p(doc, '邮箱只用于独立订阅服务和邮件发送，不进入公开源码或公开档案。各订阅人分别接收邮件，不公开其他人的邮箱。管理人员可通过受保护的后台检查登记数量和发送结果，而不在公共日志展示地址。')
h(doc, '把一次阅读变成持续追踪')
p(doc, '当日事件进入增量档案，既有三年历史语料继续保留。按国家、地区、主题、机构和时间检索，可以把连续报道还原为议题的演变过程；语料分析为专题与年度研究提供可复查的文本基础。统计反映平台收录文本的分布，不直接等同于现实事件的发生频率。')
p(doc, '建议从全球现场选择事件，阅读中文概括后按需打开原文；需要集中阅读时下载今日简报，需要持续跟踪时订阅周报，研究时再进入文本数据库和相关专题。')
url(doc, '格润网站', SITE)
url(doc, '能源技术入口', SITE+'?mode=energy')
doc.save(OUT/'格润新闻资讯与推送_对外说明.docx')

doc = document('科产融合案例库阿里云部署教程', '静态网站与独立域名  |  操作说明  |  2026年10月11日')
h(doc, '部署方案与准备事项')
p(doc, '现有案例库是静态前端与公开JSON数据，不必为页面浏览购买GPU或另行运行大模型。建议先采用“阿里云OSS静态托管＋自有域名”，保留GitHub Pages作为原站；原有GitHub Actions继续采编和导出。学生认证只用于申请优惠，不是部署前提。域名、存储、外网流量、证书或加速服务可能收费，购买前核对订单和续费价格，先设置账单预警。')
h(doc, '第一步  账号与学生验证')
p(doc, '登录自己长期使用的阿里云账号并完成实名认证。需要学生权益时，进入高校用云支持计划，点击“我是学生”下的申请入口，用同一实名的支付宝扫码并按页面授权核验。内地普通全日制专科、本科、硕士、博士在校生可按官方规则认证；具体权益以账号页面为准，不默认覆盖全部部署成本。')
h(doc, '第二步  注册域名并确定托管地域')
p(doc, '在阿里云域名注册页面查询拟用名称，核对注册价和续费价，完成持有人信息与域名实名认证。示例统一使用cases.example.com表示您实际购买域名下的子域名，示例不能直接作为真实访问地址。')
p(doc, '长期面向中国内地用户，优先评估内地托管及ICP备案路径；内地OSS绑定域名需要备案。备案主体、域名持有人及可用备案资源须按官方要求核对，不能认为购买域名就已完成备案。境外地域的接入与合规要求应分别核实；地域选定前测试目标用户的实际访问情况。')
h(doc, '第三步  创建专用OSS Bucket')
p(doc, '在控制台搜索“对象存储 OSS”，创建仅用于公开网站的Bucket，选择已确定地域、标准存储，关闭不需要的加速和额外服务。该Bucket不得混放订阅邮箱、API Key、人员原始资料、研究底稿或Git目录。只为网站公开对象授予匿名读取权限，保持匿名写入关闭；可禁止匿名列举目录。')
p(doc, '网站必须使用自定义域名。OSS默认Bucket域名访问HTML可能触发下载，不应把该地址当作最终网页链接。先完成域名与访问配置，再验收浏览器正常显示。')
doc.add_page_break()
h(doc, '第四步  获取可发布网站并上传')
p(doc, '进入GitHub仓库的Actions，打开最新成功的“每日更新气候文本数据平台”运行，在Artifacts下载github-pages。解压下载包，再解开其中的artifact.tar，确认最外层含index.html、app.js、styles.css和data文件夹。不要上传整个项目仓库，也不要把压缩包本身当作首页。')
p(doc, '在OSS文件管理中上传解包后的完整网站目录，保留相对路径。首次部署可人工上传；后续更新用同名覆盖，不先删除整个Bucket。只复制科产融合页面而漏掉共享脚本、样式、图标或数据，会导致筛选和对比失效，因此初次建议部署完整静态站点。')
p(doc, '在Bucket的静态网站设置中，将默认首页设为index.html，错误处理使用返回404或专门错误页，不把不存在的数据文件一律返回首页。HTML和JSON应采用短缓存或需重新验证的缓存策略；带版本号的静态资源可采用较长缓存。')
h(doc, '第五步  绑定域名与开启HTTPS')
p(doc, '在Bucket的域名管理中添加cases.example.com，按控制台提示完成归属验证。在域名DNS解析中添加CNAME，将主机记录cases指向控制台给出的Bucket目标地址，不手工猜测地址。内地Bucket先确认备案状态。')
p(doc, '按OSS自定义域名的证书配置要求，申请并绑定对应域名的SSL证书；如需强制HTTP跳转HTTPS，采用已配置的网关或CDN功能，并先核对费用。浏览器使用https://cases.example.com访问，确认无证书警告、混合内容或强制下载。')
h(doc, '第六步  设置案例库入口')
p(doc, '对外入口使用https://cases.example.com/?mode=energy#innovationCases。该参数进入能源技术模式，锚点定位到科产融合。首页仍保留原平台，便于读者进一步查阅新闻、报告与企业情报；若希望域名首页直接进入案例库，应在独立镜像中设置默认入口，不修改GitHub原站的气候模式。')
p(doc, '验收时至少测试案例关键词搜索、国家与中国省份筛选、案例详情、2—3例对比、国家与区域模式、CSV导出及手机端菜单。核对案例数与当前原站一致。两份推文中的网址需在新域名验收通过后再替换，不提前宣传尚未上线的地址。')
doc.add_page_break()
h(doc, '第七步  接上持续更新，而非一次性上传')
p(doc, '原站更新并不会自动同步到OSS。正式使用前应增加一个独立镜像发布任务：仅在原Pages工作流成功后，读取同一次合格github-pages产物，解包并同步到专用Bucket。不要重新调用模型或从零重采数据；这样两个站点使用同一发布快照，也不会增加每日编译次数。')
p(doc, '在RAM访问控制中建立专用部署身份，仅授予目标Bucket所需的对象写入和必要列举权限，不授予全账号管理权限。优先采用短期身份凭据；若使用AccessKey，其ID与Secret仅保存在GitHub Actions Secrets，不进入代码、本地教程、命令截图或日志。')
p(doc, '建议统一命名：ALIYUN_OSS_ACCESS_KEY_ID、ALIYUN_OSS_ACCESS_KEY_SECRET用于部署凭据；ALIYUN_OSS_BUCKET、ALIYUN_OSS_REGION、ALIYUN_SITE_URL用于仓库Variables。为镜像任务设置独立开关ALIYUN_DEPLOY_ENABLED，首次人工运行并确认成功后再启用。上述阿里云任务尚未在账号中部署，配置好地域、Bucket及域名后才可完成接入。')
p(doc, '上传时先保存旧版本，数据与脚本先上传、index.html最后上传，并在末尾写入本次快照标记；更严格的发布可采用版本化目录和入口切换。不要使用带删除参数的全目录同步覆盖不明目标。上线后检查连续两次自动任务均成功、快照标记变化且旧数据仍可检索，再视为自动镜像验收通过。')
h(doc, '订阅和问答需要单独验收')
p(doc, 'OSS只托管静态文件，不会执行邮箱登记、SMTP发信或模型问答。现有这些功能由独立服务承担，换域名不会自动迁移，也不能保证消除其他用户的网络问题。新域名上线时应将其准确Origin加入后端允许列表；只部署案例库可先不开放无关服务。')
p(doc, '如目标用户仍无法访问现有Cloudflare接口，应在可稳定访问的阿里云服务上部署受控订阅接入或迁移后端，不能仅增加同一受限域下的别名。接入必须保持HTTPS、固定上游、请求大小和频率限制、不记录邮箱请求体；订阅名单继续存于私有服务，不放OSS。周报名单读取须与公开登记入口指向同一名单。')
p(doc, '域名与Bucket准备完成后，需要再执行一次镜像发布配置和跨用户测试。本次仅提供部署方案与操作步骤，没有代购域名、开通付费资源，也没有宣称阿里云镜像已经上线。')
doc.add_page_break()
h(doc, '上线检查清单')
for text in [
    '域名实名认证、托管地域、备案及HTTPS符合所选服务要求；预算预警与续费提醒已设置。',
    '桌面、手机和微信内置浏览器均能打开科产融合入口，搜索、筛选、对比和导出可用。',
    '公开Bucket仅含发布产物；源码与产物中不存在密钥、订阅列表或未获公开授权的人员资料。',
    '镜像任务复用原站合格快照，连续两次自动发布通过；失败时保留上一版，具备回滚路径。',
    '若启用订阅，请两位此前失败的用户重新提交；页面确认登记后，管理端人数增加，再验证周报送达。不要将测试邮箱写入Git或公共日志。',
]: p(doc, '□ '+text)
h(doc, '官方参考入口')
for label, value in [
    ('学生身份验证','https://help.aliyun.com/zh/account/student-identity-verification'),
    ('高校用云支持计划','https://university.aliyun.com/'),
    ('OSS静态网站托管','https://help.aliyun.com/zh/oss/user-guide/hosting-static-websites'),
    ('OSS自定义域名','https://help.aliyun.com/zh/oss/user-guide/access-buckets-via-custom-domain-names'),
    ('ICP备案申请概述','https://help.aliyun.com/zh/icp-filing/basic-icp-service/user-guide/icp-filing-application-overview'),
]: url(doc, label, value)
p(doc, '操作入口及服务条款可能调整，执行时以本人账号控制台与上述官方文档为准。')
doc.save(OUT/'科产融合案例库_阿里云部署教程.docx')
print(json.dumps({'documents':3,'cases':stats['cases'],'models':model_count,'energy_observations':ENERGY['counts']['observations'],'output':str(OUT)},ensure_ascii=False))
